from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "data" / "source-sheets"
FILES = {
    "social_listening": SOURCE_DIR / "social-listening.xlsx",
    "competitor_followers": SOURCE_DIR / "competitor-followers.xlsx",
    "hopi_followers": SOURCE_DIR / "hopi-followers.xlsx",
}


def is_formula(value: object) -> bool:
    return isinstance(value, str) and value.startswith("=")


def compact_row(values: tuple[object, ...]) -> list[object | None]:
    output = [value for value in values]
    while output and output[-1] is None:
        output.pop()
    return output


def inspect_workbook(path: Path) -> dict[str, object]:
    workbook = load_workbook(path, data_only=False, read_only=False)
    sheets: list[dict[str, object]] = []
    for worksheet in workbook.worksheets:
        non_empty_rows: list[dict[str, object]] = []
        formula_count = 0
        for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            values = compact_row(row)
            formula_count += sum(1 for value in values if is_formula(value))
            if any(value not in (None, "") for value in values) and len(non_empty_rows) < 18:
                non_empty_rows.append({"row": row_index, "values": values})

        sheets.append(
            {
                "title": worksheet.title,
                "max_row": worksheet.max_row,
                "max_column": worksheet.max_column,
                "merged_ranges": [str(item) for item in worksheet.merged_cells.ranges],
                "formula_count": formula_count,
                "sample_non_empty_rows": non_empty_rows,
            }
        )
    return {"file": path.name, "sheet_count": len(sheets), "sheets": sheets}


def main() -> None:
    inventory = {name: inspect_workbook(path) for name, path in FILES.items()}
    output = SOURCE_DIR / "inventory.json"
    output.write_text(json.dumps(inventory, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(json.dumps({name: [sheet["title"] for sheet in data["sheets"]] for name, data in inventory.items()}, ensure_ascii=False))


if __name__ == "__main__":
    main()
